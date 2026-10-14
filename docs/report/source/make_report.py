"""Assemble the full project report (two-pass: collect headings, build, measure pages, rebuild)."""
import os, re, sys, shutil, importlib
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt, Cm
from builder import Report, TEXT_W_CM, FONT, set_cell_margins, no_borders
from refs_stub import REFS
from word_export import export_pdf

OUT_DOCX = os.path.join(os.path.dirname(HERE), "major_project_final_report - Recruitment.docx")
WORK_DOCX = os.path.join(HERE, "report_work.docx")
WORK_PDF = os.path.join(HERE, "report_work.pdf")

TITLE = "AI-Powered Intelligent Recruitment System: Explainable, Bias-Aware Candidate Ranking"
SHORT = "AI-Powered Intelligent Recruitment System"
TEAM = [("MITHILA R", "1VK23IS034"), ("SNEHA", "1VK23IS052"), ("YASHASWINI R", "1VK23IS062")]

ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("ATS", "Applicant Tracking System"),
    ("BERT", "Bidirectional Encoder Representations from Transformers"), ("BM25", "Best Matching 25 (Okapi BM25)"),
    ("CGPA", "Cumulative Grade Point Average"), ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"),
    ("CSS", "Cascading Style Sheets"), ("DFD", "Data Flow Diagram"), ("DOCX", "Office Open XML Document"),
    ("ER", "Entity-Relationship"), ("ESCO", "European Skills, Competences, Qualifications and Occupations"),
    ("EU", "European Union"), ("FA*IR", "Fair Top-k Ranking Algorithm"), ("GDPR", "General Data Protection Regulation"),
    ("GUI", "Graphical User Interface"), ("HIRE", "Hire verdict of the council"), ("HTML", "Hypertext Markup Language"),
    ("HTTP", "Hypertext Transfer Protocol"), ("IDF", "Inverse Document Frequency"), ("INR", "Indian Rupee"),
    ("JD", "Job Description"), ("JSON", "JavaScript Object Notation"), ("JSX", "JavaScript XML"),
    ("LLM", "Large Language Model"), ("LPA", "Lakhs Per Annum"), ("MiniLM", "Miniature Language Model"),
    ("ML", "Machine Learning"), ("nDCG", "Normalised Discounted Cumulative Gain"), ("NLP", "Natural Language Processing"),
    ("OCR", "Optical Character Recognition"), ("ORM", "Object-Relational Mapping"),
    ("PBKDF2", "Password-Based Key Derivation Function 2"), ("PDF", "Portable Document Format"),
    ("PK / FK", "Primary Key / Foreign Key"), ("RAG", "Retrieval-Augmented Generation"),
    ("REST", "Representational State Transfer"), ("RQ", "Research Question"), ("SBERT", "Sentence-BERT"),
    ("SHA-256", "Secure Hash Algorithm, 256-bit"), ("SMTP", "Simple Mail Transfer Protocol"), ("SQL", "Structured Query Language"),
    ("SRE", "Site Reliability Engineering"), ("TC", "Test Case"), ("TF", "Term Frequency"),
    ("TF-IDF", "Term Frequency\u2013Inverse Document Frequency"), ("UAT", "User Acceptance Testing"), ("UI", "User Interface"),
    ("URL", "Uniform Resource Locator"), ("USD", "United States Dollar"), ("VKIT", "Vivekananda Institute of Technology"),
    ("VTU", "Visvesvaraya Technological University"), ("XAI", "Explainable Artificial Intelligence"),
]

CHAPTERS = ["ch1", "ch2", "ch3", "ch4", "ch5", "ch6", "ch7", "ch8", "ch9"]


def extract_banner():
    from pptx import Presentation
    out = os.path.join(HERE, "banner.png")
    if os.path.exists(out):
        return out
    p = Presentation(r"C:\Users\Arjun\Downloads\FINAL PPT.pptx")
    for sh in p.slides[2].shapes:
        if sh.shape_type == 13 and sh.width > 11000000:
            with open(out, "wb") as f:
                f.write(sh.image.blob)
            return out
    return None


def cover(r):
    r.start_cover()
    banner = extract_banner()
    if banner:
        p = r.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(banner, width=Cm(TEXT_W_CM + 1.0))
    r.center("Department of Information Science and Engineering", size=14, bold=True, space_before=10, space_after=22)
    r.center("MAJOR PROJECT PHASE II REPORT", size=15, bold=True, space_after=6)
    r.center("on", size=12, italic=True, space_after=10)
    r.center(f"\u201c{TITLE}\u201d", size=19, bold=True, space_after=18, line=1.25)
    r.center("Submitted in partial fulfilment of the requirements for the award of the degree of", size=12, space_after=4)
    r.center("BACHELOR OF ENGINEERING", size=14, bold=True, space_after=2)
    r.center("in", size=12, italic=True, space_after=2)
    r.center("INFORMATION SCIENCE AND ENGINEERING", size=14, bold=True, space_after=18)
    r.center("Submitted by", size=12, italic=True, space_after=8)
    t = r.doc.add_table(rows=len(TEAM), cols=2)
    t.alignment = 1
    no_borders(t)
    for i, (n, usn) in enumerate(TEAM):
        for j, v in enumerate((n, usn)):
            c = t.rows[i].cells[j]
            c.width = Cm(6.0) if j == 0 else Cm(4.5)
            pp = c.paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pp.paragraph_format.space_after = Pt(4)
            rr = pp.add_run(v)
            rr.bold = True
            rr.font.size = Pt(13)
            rr.font.name = FONT
    r.center("Under the guidance of", size=12, italic=True, space_before=16, space_after=4)
    r.center("[Guide name], [Designation]", size=13, bold=True, space_after=2)
    r.center("Department of Information Science and Engineering, VKIT", size=12, space_after=14)
    r.center("Vivekananda Institute of Technology, Bengaluru", size=13, bold=True, space_after=2)
    r.center("Affiliated to Visvesvaraya Technological University, Belagavi", size=11, space_after=2)
    r.center("Academic Year 2026\u20132027", size=12, bold=True, space_before=8)


def signature_table(r, cols):
    t = r.doc.add_table(rows=2, cols=len(cols))
    no_borders(t)
    for j, (name, role) in enumerate(cols):
        c1, c2 = t.rows[0].cells[j], t.rows[1].cells[j]
        for c, txt, b in ((c1, name, True), (c2, role, False)):
            c.width = Cm(TEXT_W_CM / len(cols))
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(26 if c is c1 else 0)
            rr = p.add_run(txt)
            rr.bold = b
            rr.font.size = Pt(11)
            rr.font.name = FONT


def front_pages(r, entries):
    toc, figs, tabs = entries
    r.start_front()
    # Certificate
    r.front_title("CERTIFICATE")
    r.p(f"This is to certify that the Major Project Phase II work entitled \u201c{TITLE}\u201d is a bonafide work carried out by **MITHILA R (1VK23IS034)**, **SNEHA (1VK23IS052)** and **YASHASWINI R (1VK23IS062)** in partial fulfilment of the requirements for the award of the degree of **Bachelor of Engineering in Information Science and Engineering** of Visvesvaraya Technological University, Belagavi, during the academic year 2026\u20132027.", line=1.6)
    r.p("It is certified that all corrections and suggestions indicated for internal assessment have been incorporated in the report. The project report has been approved as it satisfies the academic requirements in respect of the project work prescribed for the said degree.", line=1.6)
    signature_table(r, [("[Guide name]", "Project Guide"), ("[HOD name]", "Head of the Department"), ("[Principal name]", "Principal")])
    # Declaration
    r.front_title("DECLARATION", new_page=True)
    r.p(f"We, **MITHILA R (1VK23IS034)**, **SNEHA (1VK23IS052)** and **YASHASWINI R (1VK23IS062)**, students of the Department of Information Science and Engineering, Vivekananda Institute of Technology, Bengaluru, hereby declare that the Major Project Phase II work entitled \u201c{TITLE}\u201d has been carried out by us under the guidance of [Guide name], Department of Information Science and Engineering, in partial fulfilment of the requirements for the award of the degree of Bachelor of Engineering of Visvesvaraya Technological University, Belagavi, during the academic year 2026\u20132027.", line=1.6)
    r.p("We further declare that this work is original, that it has not been submitted to any other university or institution for the award of any degree, and that all sources used have been acknowledged in the bibliography. AI-assisted tools were used for parts of the code generation and drafting; we have reviewed all the material and take responsibility for its content.", line=1.6)
    t = r.doc.add_table(rows=len(TEAM), cols=2)
    no_borders(t)
    for i, (n, usn) in enumerate(TEAM):
        for j, v in enumerate((n, usn)):
            c = t.rows[i].cells[j]
            c.width = Cm(8.0)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            rr = p.add_run(v)
            rr.bold = True
            rr.font.size = Pt(12)
            rr.font.name = FONT
    # Acknowledgement
    r.front_title("ACKNOWLEDGEMENT", new_page=True)
    r.p("We express our sincere gratitude to our institution, Vivekananda Institute of Technology, for providing the environment and the facilities required to complete this work.", line=1.6)
    r.p("We thank [Principal name], Principal, and [HOD name], Head of the Department of Information Science and Engineering, for their encouragement and support throughout the project.", line=1.6)
    r.p("We are deeply grateful to our project guide, [Guide name], for the valuable guidance, the constructive criticism and the time that were given to us at every stage of the work. We also thank the project coordinators and all the faculty members of the department for their suggestions during the reviews.", line=1.6)
    r.p("Finally, we thank our parents and friends for their constant support and motivation.", line=1.6)
    p = r.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(18)
    rr = p.add_run("MITHILA R\nSNEHA\nYASHASWINI R")
    rr.bold = True
    rr.font.name = FONT
    # Abstract
    r.front_title("ABSTRACT", new_page=True)
    r.p("Keyword-based applicant tracking systems rank candidates by the overlap between a resume and a job description. They reward vocabulary rather than evidence, they cannot relate similar skills, and they give no explanation of their decisions, while recent audits show that neural resume screeners can inherit demographic bias. This project presents the AI-Powered Intelligent Recruitment System, an explainable, bias-aware and offline-capable screening system that treats every skill as a claim that needs evidence. A resume parser with an optical character recognition fallback produces structured profiles, and a 79-skill ontology with alias normalisation, fuzzy matching and partial credit reports matched, partially matched and missing skills. A seven-factor hybrid ranker combines skill match, semantic similarity, experience, education, projects, certifications and resume quality, and an evidence-weighted skill coverage gives half credit to skills that are listed but not demonstrated.", line=1.4, size=11.5)
    r.p("A blind-screening filter removes names, contact details and gendered words before any evaluation, and a council of four agents (Tech Lead, Culture and Retention, Fair-Hiring Auditor, and Seniority and Compensation) produces a consensus verdict with knock-out gates, red flags, confidence, dissent and human-review triggers. The system also contains a GitHub authenticity inspector, a hybrid dense and BM25 talent search, an interview-question generator and a retrieval-grounded recruiter chatbot. It is implemented as a Python core, a FastAPI service with a relational database, a Streamlit workspace and a React dashboard, and it is deployed as a multi-user web service with server-side accounts (salted PBKDF2 hashing, expiring sessions, verified Google sign-in) and a MongoDB Atlas database. A suite of 78 automated tests passes. On a pilot of twelve synthetic resumes, a screening takes about 30 milliseconds per resume, a keyword-stuffed resume that a keyword baseline scores like a genuine candidate is ranked tenth and rejected, scores did not change in 96 name-substitution runs, and rankings were stable under random weight changes of 20 percent. All data are synthetic, and the limitations of the evaluation are stated explicitly.", line=1.4, size=11.5)
    p = r.p("**Keywords:** Resume Screening, Applicant Tracking System, Explainable AI, Fairness, Skill Ontology, Hybrid Retrieval, Multi-Agent Systems, Blind Screening.", indent=False, line=1.4, size=11.5, align="left")
    p.paragraph_format.space_before = Pt(10)
    # TOC
    r.front_title("TABLE OF CONTENTS", new_page=True)
    r.toc_lines(toc, "toc")
    r.front_title("LIST OF FIGURES", new_page=True)
    r.toc_lines(figs, "figs")
    r.front_title("LIST OF TABLES", new_page=True)
    r.toc_lines(tabs, "tabs")
    r.front_title("LIST OF ABBREVIATIONS", new_page=True)
    r.abbreviations(ABBR)


def build_doc(pages, entries):
    r = Report(SHORT, pages=pages)
    cover(r)
    front_pages(r, entries)
    for name in CHAPTERS:
        mod = importlib.import_module(name)
        importlib.reload(mod)
        mod.build(r)
    r.bibliography(REFS)
    return r


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def measure(pdf_path, r):
    import fitz
    d = fitz.open(pdf_path)
    texts = [norm(p.get_text()) for p in d]
    body0 = next(i for i, t in enumerate(texts) if "CHAPTER 1" in t and "INTRODUCTION" in t and "Dept. of ISE, VKIT" in t)
    pages = {}
    cur = body0

    def find(pattern_fn, start):
        for i in range(start, len(texts)):
            if pattern_fn(texts[i]):
                return i
        return None

    for level, num, text in r.toc:
        if level == 1 and num:
            i = find(lambda t: f"CHAPTER {num} {text}" in t, cur)
        elif level == 1:
            i = find(lambda t: text in t and "Dept. of ISE" in t, cur)
        else:
            i = find(lambda t: norm(f"{num} {text}") in t, cur)
        if i is None:
            print("  !! not found:", num, text)
            continue
        cur = i
        pages[("toc", num, text)] = i - body0 + 1
    for kind, lst in (("figs", r.figs), ("tabs", r.tabs)):
        cur2 = body0
        for label, caption in lst:
            plain = norm(re.sub(r"\*\*|\*|`|~|\^|\[@[^\]]*\]", "", caption))
            key = norm(f"{label}: {plain[:40]}")
            i = find(lambda t: key in t, cur2)
            if i is None:
                i = find(lambda t: norm(f"{label}:") in t, cur2)
            if i is None:
                print("  !! not found:", label)
                continue
            cur2 = i
            pages[(kind, label)] = i - body0 + 1
    return pages, body0, len(texts)


def main():
    # pass 0: collect entries
    r0 = build_doc({}, ([], [], []))
    entries = (list(r0.toc), list(r0.figs), list(r0.tabs))
    print("entries:", len(entries[0]), "toc,", len(entries[1]), "figs,", len(entries[2]), "tables; cited refs:", len(r0.cite_order))
    unused = [k for k in REFS if k not in r0.cites]
    print("UNCITED:", unused)
    # pass 1: placeholder page numbers
    r1 = build_doc({}, entries)
    r1.save(WORK_DOCX)
    n = export_pdf(WORK_DOCX, WORK_PDF)
    pages, body0, total = measure(WORK_PDF, r1)
    print("pass1 pages:", n, "body starts at pdf page", body0 + 1)
    # pass 2: real page numbers
    r2 = build_doc(pages, entries)
    r2.save(WORK_DOCX)
    n2 = export_pdf(WORK_DOCX, WORK_PDF)
    pages2, body02, total2 = measure(WORK_PDF, r2)
    stable = pages2 == pages
    print("pass2 pages:", n2, "stable:", stable)
    if not stable:
        r3 = build_doc(pages2, entries)
        r3.save(WORK_DOCX)
        n3 = export_pdf(WORK_DOCX, WORK_PDF)
        print("pass3 pages:", n3)
    shutil.copy(WORK_DOCX, OUT_DOCX)
    print("saved", OUT_DOCX)


if __name__ == "__main__":
    main()
