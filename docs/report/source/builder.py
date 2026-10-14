"""Report builder: python-docx based, with native equations (OMML), tables, figures and a two-pass TOC."""
import copy, os, re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from lxml import etree
import latex2mathml.converter as l2m

FONT = "Times New Roman"
XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
HERE = os.path.dirname(os.path.abspath(__file__))

TEXT_W_CM = 21.0 - 3.5 - 2.5  # A4 text width


def latex_to_omml(latex):
    mml = l2m.convert(latex)
    return XSL(etree.fromstring(mml.encode())).getroot()


def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_cell_margins(cell, top=40, bottom=40, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def table_borders(table, sz=4, color="000000", inner=True):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right") + (("insideH", "insideV") if inner else ()):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def keep_with_next(p, on=True):
    p.paragraph_format.keep_with_next = on


def add_field(run, instr):
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "1"
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        run._r.append(e)


def para_border(p, side="bottom", sz=8, color="000000", space=1):
    pPr = p._p.get_or_add_pPr()
    b = pPr.find(qn("w:pBdr"))
    if b is None:
        b = OxmlElement("w:pBdr")
        pPr.append(b)
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    b.append(e)


class Report:
    def __init__(self, short_title, pages=None):
        self.doc = Document()
        self.short_title = short_title
        self.pages = pages or {}
        self.cites = {}
        self.cite_order = []
        self.toc = []       # (level, number, text)
        self.figs = []      # (label, caption)
        self.tabs = []      # (label, caption)
        self.chapter_no = 0
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self.code_n = 0
        self._styles()
        self._first_section_done = False

    # ------------------------------------------------------------------ styles & page
    def _styles(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = FONT
        n.font.size = Pt(12)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        n.paragraph_format.space_after = Pt(0)
        n.paragraph_format.space_before = Pt(0)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.3)
        sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.0)
        cp = self.doc.core_properties
        cp.title = "AI-Powered Intelligent Recruitment System: Explainable, Bias-Aware Candidate Ranking"
        cp.author = "Mithila R, Sneha, Yashaswini R"
        cp.subject = "Major Project Phase II Report, Dept. of ISE, VKIT"

    def _setup_section(self, sec, borders=False, pgfmt=None, start=None, titlepg=False):
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.3)
        sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.0)
        sectPr = sec._sectPr
        for old in sectPr.findall(qn("w:pgBorders")) + sectPr.findall(qn("w:pgNumType")):
            sectPr.remove(old)
        if borders:
            pb = OxmlElement("w:pgBorders")
            pb.set(qn("w:offsetFrom"), "page")
            for side in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:val"), "double")
                e.set(qn("w:sz"), "6")
                e.set(qn("w:space"), "24")
                e.set(qn("w:color"), "000000")
                pb.append(e)
            # pgBorders must come after pgMar in schema order
            pgMar = sectPr.find(qn("w:pgMar"))
            pgMar.addnext(pb)
        if pgfmt:
            pn = OxmlElement("w:pgNumType")
            pn.set(qn("w:fmt"), pgfmt)
            if start is not None:
                pn.set(qn("w:start"), str(start))
            anchor = sectPr.find(qn("w:pgBorders"))
            if anchor is None:
                anchor = sectPr.find(qn("w:pgMar"))
            anchor.addnext(pn)
        sec.different_first_page_header_footer = titlepg

    def _clear(self, hf):
        hf.is_linked_to_previous = False
        for p in hf.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        p = hf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        return p

    def _footer_roman(self, sec):
        f = self._clear(sec.footer)
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = f.add_run()
        r.font.size = Pt(11)
        r.font.name = FONT
        add_field(r, "PAGE")
        sec.header.is_linked_to_previous = False
        self._clear(sec.header)

    def _footer_body(self, ftr):
        f = self._clear(ftr)
        f.paragraph_format.tab_stops.clear_all() if hasattr(f.paragraph_format.tab_stops, "clear_all") else None
        ts = f.paragraph_format.tab_stops
        ts.add_tab_stop(Cm(TEXT_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        para_border(f, "top", sz=12, space=4)
        for txt, bold in (("Dept. of ISE, VKIT", False), ("\t2026\u20132027", False), ("\t", False)):
            r = f.add_run(txt)
            r.font.size = Pt(10.5)
            r.font.name = FONT
        r = f.add_run()
        r.font.size = Pt(10.5)
        r.font.name = FONT
        add_field(r, "PAGE")

    def _header_body(self, hdr):
        h = self._clear(hdr)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = h.add_run(self.short_title)
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = FONT
        para_border(h, "bottom", sz=12, space=2)

    # ------------------------------------------------------------------ section creators
    def start_cover(self):
        sec = self.doc.sections[0]
        self._setup_section(sec)
        sec.footer.is_linked_to_previous = False
        self._clear(sec.footer)
        sec.header.is_linked_to_previous = False
        self._clear(sec.header)

    def start_front(self, borders=True, start=1):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_section(sec, borders=borders, pgfmt="lowerRoman", start=start)
        self._footer_roman(sec)
        return sec

    def new_front_page(self):
        self.doc.add_page_break()

    def start_body_section(self, restart=False):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_section(sec, borders=False, pgfmt="decimal", start=1 if restart else None, titlepg=True)
        # normal header/footer
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        self._header_body(sec.header)
        self._footer_body(sec.footer)
        # first page: no header, footer kept
        sec.first_page_header.is_linked_to_previous = False
        self._clear(sec.first_page_header)
        sec.first_page_footer.is_linked_to_previous = False
        self._footer_body(sec.first_page_footer)
        return sec

    # ------------------------------------------------------------------ inline markup
    TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\[@[^\]]+\]|~[^~]+~|\^[^^]+\^)")

    def cite_text(self, keys):
        nums = []
        for k in keys:
            k = k.strip()
            if k not in self.cites:
                self.cites[k] = len(self.cites) + 1
                self.cite_order.append(k)
            nums.append(self.cites[k])
        return ", ".join(f"[{n}]" for n in nums)

    def add_runs(self, p, text, size=None, bold=None, italic=None, color=None, font=None):
        for tok in self.TOKEN.split(text):
            if not tok:
                continue
            r = None
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                r = p.add_run(tok[2:-2])
                r.bold = True
            elif tok.startswith("[@"):
                keys = [k.strip().lstrip("@") for k in tok[2:-1].split(",")]
                r = p.add_run(self.cite_text(keys))
            elif tok.startswith("`") and tok.endswith("`"):
                r = p.add_run(tok[1:-1])
                r.font.name = "Consolas"
                r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt((size or 12) - 1.5)
                continue
            elif tok.startswith("~") and tok.endswith("~") and len(tok) > 2:
                r = p.add_run(tok[1:-1])
                r.font.subscript = True
            elif tok.startswith("^") and tok.endswith("^") and len(tok) > 2:
                r = p.add_run(tok[1:-1])
                r.font.superscript = True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                r = p.add_run(tok[1:-1])
                r.italic = True
            else:
                r = p.add_run(tok)
            if r is not None:
                r.font.name = font or FONT
                if size:
                    r.font.size = Pt(size)
                if bold:
                    r.bold = True
                if italic:
                    r.italic = True
                if color:
                    r.font.color.rgb = color

    # ------------------------------------------------------------------ text blocks
    def p(self, text, indent=True, align="justify", size=12, space_after=6, line=1.5, left=0.0, keep=False, bold=None, italic=None):
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        pf.line_spacing = line
        pf.space_after = Pt(space_after)
        if indent:
            pf.first_line_indent = Cm(0.8)
        if left:
            pf.left_indent = Cm(left)
        pf.widow_control = True
        if keep:
            pf.keep_with_next = True
        self.add_runs(para, text, size=size, bold=bold, italic=italic)
        return para

    def bullets(self, items, size=12, marker="\u2022", left=1.0, hang=0.5, space_after=3, line=1.3):
        for it in items:
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(left + hang)
            pf.first_line_indent = Cm(-hang)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = line
            pf.space_after = Pt(space_after)
            pf.tab_stops.add_tab_stop(Cm(left + hang))
            self.add_runs(para, f"{marker}\t{it}", size=size)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def numbered(self, items, size=12, left=1.0, hang=0.7, space_after=3, line=1.3):
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(left + hang)
            pf.first_line_indent = Cm(-hang)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = line
            pf.space_after = Pt(space_after)
            pf.tab_stops.add_tab_stop(Cm(left + hang))
            self.add_runs(para, f"{i}.\t{it}", size=size)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def chapter(self, title):
        self.chapter_no += 1
        self.fig_n = self.tab_n = self.eq_n = self.code_n = 0
        self.start_body_section(restart=(self.chapter_no == 1))
        p1 = self.doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(20)
        p1.paragraph_format.space_after = Pt(10)
        r = p1.add_run(f"CHAPTER {self.chapter_no}")
        r.bold = True
        r.font.size = Pt(14)
        r.font.name = FONT
        p2 = self.doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(16)
        r = p2.add_run(title.upper())
        r.bold = True
        r.font.size = Pt(18)
        r.font.name = FONT
        keep_with_next(p1)
        keep_with_next(p2)
        self.toc.append((1, str(self.chapter_no), title.upper()))

    def h1(self, num, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(14)
        pf.space_after = Pt(8)
        pf.keep_with_next = True
        pf.left_indent = Cm(1.2)
        pf.first_line_indent = Cm(-1.2)
        pf.tab_stops.add_tab_stop(Cm(1.2))
        r = p.add_run(f"{num}\t{title}")
        r.bold = True
        r.font.size = Pt(13)
        r.font.name = FONT
        self.toc.append((2, num, title))

    def h2(self, num, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.left_indent = Cm(1.5)
        pf.first_line_indent = Cm(-1.5)
        pf.tab_stops.add_tab_stop(Cm(1.5))
        r = p.add_run(f"{num}\t{title}")
        r.bold = True
        r.font.size = Pt(12)
        r.font.name = FONT

    def theme(self, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(14)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        r = p.add_run(title)
        r.bold = True
        r.underline = True
        r.font.size = Pt(12.5)
        r.font.name = FONT

    def h3(self, title):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(8)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(12)
        r.font.name = FONT

    # ------------------------------------------------------------------ equation
    def eq(self, latex, label=True):
        if label:
            self.eq_n += 1
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = Pt(4)
        pf.space_after = Pt(8)
        pf.line_spacing = 1.15
        pf.keep_together = True
        ts = pf.tab_stops
        ts.add_tab_stop(Cm(TEXT_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        r0 = para.add_run("\t")
        omml = latex_to_omml(latex)
        para._p.append(omml)
        if label:
            r = para.add_run(f"\t({self.chapter_no}.{self.eq_n})")
            r.font.name = FONT
            r.font.size = Pt(12)
        return para

    def where(self, items, intro="where"):
        self.p(intro + ":", indent=False, space_after=2, line=1.3, keep=True)
        for sym, desc in items:
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(2.4)
            pf.first_line_indent = Cm(-1.4)
            pf.tab_stops.add_tab_stop(Cm(2.4))
            pf.space_after = Pt(2)
            pf.line_spacing = 1.25
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.add_runs(para, sym, size=11.5, italic=True)
            self.add_runs(para, "\t" + desc, size=11.5)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    # ------------------------------------------------------------------ figures, tables, code
    def fig(self, path, caption, width_cm=13.5, frame=False):
        self.fig_n += 1
        label = f"Figure {self.chapter_no}.{self.fig_n}"
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.keep_with_next = True
        run = para.add_run()
        run.add_picture(path, width=Cm(width_cm))
        if frame:
            # thin border around screenshot
            pic = run._r.xpath(".//pic:pic")[0]
            spPr = pic.find(qn("pic:spPr"))
            ln = parse_xml(f'<a:ln {nsdecls("a")} w="6350"><a:solidFill><a:srgbClr val="999999"/></a:solidFill></a:ln>')
            spPr.append(ln)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.line_spacing = 1.1
        r = cap.add_run(f"{label}: ")
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = FONT
        self.add_runs(cap, caption, size=11)
        self.figs.append((label, caption))
        return label

    def ref_next_fig(self):
        return f"Figure {self.chapter_no}.{self.fig_n + 1}"

    def ref_next_tab(self):
        return f"Table {self.chapter_no}.{self.tab_n + 1}"

    def ref_next_eq(self):
        return f"({self.chapter_no}.{self.eq_n + 1})"

    def table(self, headers, rows, caption, colw, size=10.5, align=None, hdr_fill="D9E2F3", bold_first=False, cap_above=True):
        self.tab_n += 1
        label = f"Table {self.chapter_no}.{self.tab_n}"
        if cap_above:
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_before = Pt(8)
            cap.paragraph_format.space_after = Pt(4)
            cap.paragraph_format.keep_with_next = True
            r = cap.add_run(f"{label}: ")
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = FONT
            self.add_runs(cap, caption, size=11)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_borders(t, sz=4)
        total = sum(colw)
        widths = [Cm(TEXT_W_CM * w / total) for w in colw]
        align = align or ["left"] * len(headers)
        for ci, h in enumerate(headers):
            c = t.rows[0].cells[ci]
            c.width = widths[ci]
            set_cell_shading(c, hdr_fill)
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.05
            self.add_runs(p, h, size=size, bold=True)
        # repeat header row
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for ri, row in enumerate(rows, start=1):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            trPr.append(cs)
            for ci, val in enumerate(row):
                c = tr.cells[ci]
                c.width = widths[ci]
                set_cell_margins(c)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                lines = str(val).split("\n")
                for li, line in enumerate(lines):
                    p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align[ci]]
                    p.paragraph_format.line_spacing = 1.05
                    self.add_runs(p, line, size=size, bold=(bold_first and ci == 0) or None)
        if len(rows) <= 14:
            for tr in t.rows[:-1]:
                for c in tr.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.keep_with_next = True
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        self.tabs.append((label, caption))
        return label

    def code(self, text, caption, size=8.5):
        self.code_n += 1
        label = f"Listing {self.chapter_no}.{self.code_n}"
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.keep_with_next = True
        r = cap.add_run(f"{label}: ")
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = FONT
        self.add_runs(cap, caption, size=11)
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t, sz=4, color="808080", inner=False)
        c = t.rows[0].cells[0]
        c.width = Cm(TEXT_W_CM)
        set_cell_shading(c, "F5F5F5")
        set_cell_margins(c, 60, 60, 110, 110)
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            p = c.paragraphs[0] if i == 0 else c.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(size)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        return label

    # ------------------------------------------------------------------ front-matter helpers
    def front_title(self, text, size=16, space_before=12, new_page=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = new_page
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)
        r.font.name = FONT
        return p

    def page_break(self):
        self.doc.add_page_break()

    def center(self, text, size=12, bold=False, italic=False, space_after=6, space_before=0, color=None, line=1.15):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = line
        self.add_runs(p, text, size=size, bold=bold, italic=italic, color=color)
        return p

    def toc_lines(self, entries, kind="toc"):
        """entries: list of (level, number, text) for TOC; or (label, caption) for lists."""
        pn_key = lambda n: str(self.pages.get(n, "00"))
        right = Cm(TEXT_W_CM)
        if kind == "toc":
            hdr = self.doc.add_paragraph()
            ts = hdr.paragraph_format.tab_stops
            ts.add_tab_stop(Cm(TEXT_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
            ts.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
            hdr.paragraph_format.space_after = Pt(6)
            r = hdr.add_run("CHAPTER NO.\tTITLE\tPAGE NO.")
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = FONT
            for level, num, text in entries:
                p = self.doc.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(2 if level == 2 else 3)
                pf.space_before = Pt(5 if level == 1 else 0)
                pf.keep_with_next = False
                if level == 1:
                    pf.tab_stops.add_tab_stop(Cm(1.2))
                    pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
                    r = p.add_run(f"{num}\t{text}\t{pn_key(('toc', num, text))}")
                    r.bold = True
                else:
                    pf.left_indent = Cm(1.2)
                    pf.tab_stops.add_tab_stop(Cm(2.4))
                    pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
                    r = p.add_run(f"{num}\t{text}\t{pn_key(('toc', num, text))}")
                r.font.size = Pt(11)
                r.font.name = FONT
        else:
            title = "FIGURE NO." if kind == "figs" else "TABLE NO."
            hdr = self.doc.add_paragraph()
            ts = hdr.paragraph_format.tab_stops
            ts.add_tab_stop(Cm(2.7), WD_TAB_ALIGNMENT.LEFT)
            ts.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
            hdr.paragraph_format.space_after = Pt(8)
            r = hdr.add_run(f"{title}\tTITLE\tPAGE NO.")
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = FONT
            for label, caption in entries:
                p = self.doc.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(5)
                pf.left_indent = Cm(2.7)
                pf.first_line_indent = Cm(-2.7)
                pf.tab_stops.add_tab_stop(Cm(2.7))
                pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
                num = label.split()[-1]
                plain = re.sub(r"\*\*|\*|`|~|\^|\[@[^\]]*\]", "", caption)
                r = p.add_run(f"{num}\t{plain}\t{pn_key((kind, label))}")
                r.font.size = Pt(11)
                r.font.name = FONT

    def abbreviations(self, items):
        t = self.doc.add_table(rows=1 + len(items), cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        no_borders(t)
        widths = [Cm(2.0), Cm(3.7), Cm(9.3)]
        for ci, h in enumerate(("SL. NO.", "ABBREVIATION", "FULL FORM")):
            c = t.rows[0].cells[ci]
            c.width = widths[ci]
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci == 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = FONT
        for i, (a, f) in enumerate(items, 1):
            row = t.rows[i]
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            trPr.append(cs)
            for ci, v in enumerate((str(i), a, f)):
                c = row.cells[ci]
                c.width = widths[ci]
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run(v)
                r.font.size = Pt(11)
                r.font.name = FONT
        # repeat header
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)

    def bibliography(self, refs):
        self.start_body_section(restart=False)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run("BIBLIOGRAPHY")
        r.bold = True
        r.font.size = Pt(16)
        r.font.name = FONT
        self.toc.append((1, "", "BIBLIOGRAPHY"))
        for i, key in enumerate(self.cite_order, 1):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(1.0)
            pf.first_line_indent = Cm(-1.0)
            pf.tab_stops.add_tab_stop(Cm(1.0))
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = 1.15
            pf.space_after = Pt(7)
            r = para.add_run(f"[{i}]\t")
            r.font.size = Pt(11)
            r.font.name = FONT
            self.add_runs(para, refs[key], size=11)

    def save(self, path):
        self.doc.save(path)
