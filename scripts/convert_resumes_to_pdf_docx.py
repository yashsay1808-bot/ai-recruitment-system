"""
Utility to convert synthetic text resumes into real PDF and DOCX files
for testing multi-format resume intake in the UI.
"""

from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import textwrap

def text_to_docx(txt_path: Path, docx_path: Path):
    doc = Document()
    with open(txt_path, "r", encoding="utf-8") as f:
        for line in f:
            line_str = line.strip()
            if line_str.startswith("NAME:") or line_str.startswith("SUMMARY:"):
                doc.add_heading(line_str, level=1)
            elif line_str.endswith(":") and len(line_str) < 30:
                doc.add_heading(line_str, level=2)
            else:
                doc.add_paragraph(line_str)
    doc.save(str(docx_path))

def text_to_pdf(txt_path: Path, pdf_path: Path):
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    width, height = letter
    y = height - 50
    line_height = 14
    margin = 50
    max_chars_per_line = 85

    with open(txt_path, "r", encoding="utf-8") as f:
        for line in f:
            wrapped = textwrap.wrap(line.strip(), width=max_chars_per_line) or [""]
            for w in wrapped:
                if y < 50:
                    c.showPage()
                    y = height - 50
                c.drawString(margin, y, w)
                y -= line_height
    c.save()

def main():
    source_dir = Path("data/synthetic_resumes_100")
    txt_files = list(source_dir.glob("*.txt"))
    if not txt_files:
        print("No text files found in", source_dir)
        return

    out_pdf_dir = source_dir / "pdf_format"
    out_docx_dir = source_dir / "docx_format"
    out_pdf_dir.mkdir(parents=True, exist_ok=True)
    out_docx_dir.mkdir(parents=True, exist_ok=True)

    print(f"Converting first 15 sample resumes to PDF and DOCX for local file upload testing...")
    for txt_file in txt_files[:15]:
        stem = txt_file.stem
        pdf_target = out_pdf_dir / f"{stem}.pdf"
        docx_target = out_docx_dir / f"{stem}.docx"

        text_to_docx(txt_file, docx_target)
        text_to_pdf(txt_file, pdf_target)
        print(f"  [OK] Converted: {stem} -> .pdf & .docx")

    print(f"\nDone! PDF files saved to: {out_pdf_dir}")
    print(f"DOCX files saved to: {out_docx_dir}")

if __name__ == "__main__":
    main()
